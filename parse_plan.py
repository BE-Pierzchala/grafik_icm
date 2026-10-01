"""Read the IO timetable workbook and turn one semester's column block into events."""
from __future__ import annotations

import datetime as dt
import hashlib
import re
from dataclasses import dataclass, field

import openpyxl

# Google Calendar event colours (colorId -> hex), from the Calendar API colors endpoint.
GOOGLE_COLORS = {
    "1": "#a4bdfc", "2": "#7ae7bf", "3": "#dbadff", "4": "#ff887c",
    "5": "#fbd75b", "6": "#ffb878", "7": "#46d6db", "8": "#e1e1e1",
    "9": "#5484ed", "10": "#51b749", "11": "#dc2127",
}
NOTICE_KEY = "__notice__"


@dataclass
class Event:
    code: str                      # subject code as in the plan, or NOTICE_KEY
    title: str
    date: dt.date
    start: dt.time | None          # None -> all-day
    end: dt.time | None
    end_date: dt.date | None = None  # for multi-day all-day events (inclusive)
    lecturer: str = ""
    room: str = ""
    row: int = 0
    excel_rgb: str | None = None
    color_id: str | None = None
    notes: list[str] = field(default_factory=list)
    semester: int = 0

    @property
    def event_id(self) -> str:
        # Deterministic id (hex is valid base32hex) -> re-runs update instead of duplicating.
        key = f"{self.semester}|{self.code}|{self.date}|{self.start}|{self.end}|{self.row}"
        return "io" + hashlib.sha1(key.encode()).hexdigest()


class PlanParser:
    def __init__(self, cfg: dict, excel_path: str):
        self.cfg = cfg
        self.wb = openpyxl.load_workbook(excel_path)
        self.ws = self.wb[cfg["plan_sheet"]]
        self.warnings: list[str] = []
        self.skipped: list[str] = []
        self._merge_map = {}
        for rng in self.ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self._merge_map[(r, c)] = rng

    # ---------- helpers ----------
    def cell(self, r: int, c: int):
        """Value of a cell, resolving merged cells to their top-left value."""
        rng = self._merge_map.get((r, c))
        if rng is not None:
            return self.ws.cell(rng.min_row, rng.min_col).value
        return self.ws.cell(r, c).value

    @staticmethod
    def _as_time(v):
        if isinstance(v, dt.time):
            return v
        if isinstance(v, dt.datetime):
            return v.time()
        if isinstance(v, str) and re.fullmatch(r"\d{1,2}:\d{2}", v.strip()):
            h, m = v.strip().split(":")
            return dt.time(int(h), int(m))
        return None

    @staticmethod
    def _as_date(v):
        if isinstance(v, dt.datetime):
            return v.date()
        if isinstance(v, dt.date):
            return v
        return None

    def _rgb(self, r, c):
        fill = self.ws.cell(r, c).fill
        rgb = getattr(fill.fgColor, "rgb", None) if fill and fill.fill_type else None
        if isinstance(rgb, str) and len(rgb) == 8 and rgb != "00000000":
            return "#" + rgb[2:].lower()
        return None

    # ---------- legend ----------
    def load_legend(self) -> dict[str, str | None]:
        ws = self.wb[self.cfg["legend_sheet"]]
        legend: dict[str, str | None] = {}
        for row in ws.iter_rows(min_row=1, values_only=True):
            a = row[0] if row else None
            b = row[1] if len(row) > 1 else None
            if not isinstance(a, str) or not a.strip():
                continue
            a = a.strip()
            if a.lower() == "skrót" or a.upper().startswith(("SEMESTR", "PRZEDMIOT")):
                continue
            # "ATSa-30 - j. ang." -> "ATSa-30", "POWI *" -> "POWI"
            code = re.split(r"\s+-\s+|\s*\*", a)[0].strip()
            title = b.strip() if isinstance(b, str) and b.strip() else None
            if title:
                title = re.sub(r"\s*\*\s*$", "", title)
            legend[code] = title
        return legend

    def resolve_title(self, code: str, legend: dict) -> str:
        override = self.cfg.get("titles", {}).get(code)
        if override:
            return override
        candidates = [code]
        if code.endswith("a"):
            candidates.append(code[:-1])          # PRa -> PR, POWIa -> POWI
        for cand in candidates:
            if cand in legend:
                title = legend[cand]
                if cand != code:
                    self.warnings.append(f"'{code}' matched legend entry '{cand}' (name: {title}).")
                if not title:
                    self.warnings.append(f"'{code}': legend has no full name -> using the code. "
                                         f"Add it under [titles] in config.toml.")
                    return code
                if cand in title:  # case-sensitive: e.g. 'Podstawy CYBa-30erbezpieczeństwa'
                    self.warnings.append(f"'{code}': legend name looks corrupted ('{title}'). "
                                         f"Override it under [titles] in config.toml.")
                return title
        self.warnings.append(f"'{code}': not found in legend -> using the code.")
        return code

    # ---------- main parse ----------
    def parse(self) -> list[Event]:
        legend = self.load_legend()
        sem_label = self.cfg["semester_header"].strip().upper()
        exclude = set(self.cfg.get("exclude", []))
        fixes = {int(k): dt.date.fromisoformat(v) for k, v in self.cfg.get("date_fixes", {}).items()}

        events: list[Event] = []
        title_cache: dict[str, str] = {}
        seen_notices = set()
        col = None                  # first column of the semester block
        week_name = None
        week_first: dt.date | None = None
        week_last: dt.date | None = None
        anomalous_days: dict[int, str] = {}

        for r in range(1, self.ws.max_row + 1):
            a = self.ws.cell(r, 1).value
            if isinstance(a, str) and a.strip().upper().startswith("TYDZIEŃ"):
                week_name, week_first, week_last = a.strip(), None, None
                col = None
                for c in range(1, self.ws.max_column + 1):
                    v = self.ws.cell(r, c).value
                    if isinstance(v, str) and v.strip().upper() == sem_label:
                        col = c
                if col is None:
                    self.warnings.append(f"Row {r}: '{sem_label}' not found in header of {week_name}.")
                continue
            if col is None:
                continue

            date = self._as_date(self.cell(r, 1))
            if date is None:
                continue
            day_rng = self._merge_map.get((r, 1))
            day_row = day_rng.min_row if day_rng else r

            # --- date sanity check within the week block ---
            if day_row in fixes:
                date = fixes[day_row]
            elif day_row not in anomalous_days:
                bad = None
                if week_last and date < week_last:
                    bad = f"earlier than previous day ({week_last})"
                elif week_first and (date - week_first).days > 6:
                    bad = f"more than 6 days after the week start ({week_first})"
                if bad:
                    anomalous_days[day_row] = f"Row {day_row} ({week_name}): date {date} is {bad}"
                else:
                    week_first = week_first or date
                    week_last = date
            if day_row in anomalous_days:
                if self.cell(r, col):
                    self.skipped.append(f"{anomalous_days[day_row]} -> skipped "
                                        f"'{self.cell(r, col)}' {self.cell(r, 2)}")
                continue

            start, end = self._as_time(self.cell(r, 2)), self._as_time(self.cell(r, 3))
            raw = self.cell(r, col)
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                continue
            raw = str(raw).strip()

            # --- notices: merged across several semester blocks (either direction,
            #     so the last block, SEMESTR 3, is handled too) ---
            rng = self._merge_map.get((r, col))
            if rng is not None and (rng.min_col < col or rng.max_col > col + 2):
                if (rng.min_row, rng.min_col) in seen_notices:
                    continue
                seen_notices.add((rng.min_row, rng.min_col))
                if not self.cfg.get("include_notices", True):
                    continue
                events.append(self._notice(raw, rng, date))
                continue

            # --- regular class ---
            # "PROa-do 17:45", "SEM3a do 12:15", "X - od 10:00"
            m = re.match(r"^(.*?)\s*(?:-\s*|\s+)(od|do)\s+(\d{1,2}:\d{2})\s*$", raw)
            code, note = raw, []
            if m:
                code = m.group(1)
                t = self._as_time(m.group(3))
                if m.group(2) == "do":
                    end = t
                else:
                    start = t
                note.append(f"Uwaga w planie: '{raw}'")
            if code in exclude:
                continue
            if start is None or end is None:
                self.warnings.append(f"Row {r}: '{raw}' on {date} has no time -> skipped.")
                continue
            if code not in title_cache:
                title_cache[code] = self.resolve_title(code, legend)
            events.append(Event(
                code=code, title=title_cache[code], date=date, start=start, end=end,
                lecturer=str(self.cell(r, col + 1) or "").strip(),
                room=str(self.cell(r, col + 2) or "").strip(),
                row=r, excel_rgb=self._rgb(r, col), notes=note,
            ))

        for e in events:
            e.semester = self.cfg["semester"]
        self._assign_colors(events)
        return events

    def _notice(self, text: str, rng, date: dt.date) -> Event:
        first_day = self._as_date(self.cell(rng.min_row, 1))
        last_day = self._as_date(self.cell(rng.max_row, 1))
        # whole day(s) if the notice covers every slot of its first and last day
        day_rng = self._merge_map.get((rng.max_row, 1))
        first_day_rng = self._merge_map.get((rng.min_row, 1))
        covers_start = first_day_rng is None or rng.min_row <= first_day_rng.min_row
        covers_end = day_rng is None or rng.max_row >= day_rng.max_row
        title = " | ".join(l.strip() for l in text.splitlines() if l.strip())[:150]
        if first_day != last_day or (covers_start and covers_end):
            return Event(code=NOTICE_KEY, title=title, date=first_day, start=None, end=None,
                         end_date=last_day, row=rng.min_row, notes=[text])
        return Event(code=NOTICE_KEY, title=title, date=first_day,
                     start=self._as_time(self.cell(rng.min_row, 2)),
                     end=self._as_time(self.cell(rng.max_row, 3)), row=rng.min_row, notes=[text])

    # ---------- colours ----------
    def _assign_colors(self, events: list[Event]):
        """Distinct Google colour per subject. Subjects with the most events are assigned
        first, so if there are more than 11 subjects the rarest ones share a colour.
        Notices get no colorId (they use the calendar's default colour)."""
        overrides = {k: str(v) for k, v in self.cfg.get("colors", {}).items()}
        counts, rgb_of = {}, {}
        for e in events:
            if e.code == NOTICE_KEY:
                continue
            counts[e.code] = counts.get(e.code, 0) + 1
            rgb_of.setdefault(e.code, e.excel_rgb)
        assigned = dict(overrides)
        used = set(overrides.values())
        strategy = self.cfg.get("color_strategy", "excel")
        for code in sorted(counts, key=lambda c: -counts[c]):
            if code in assigned:
                continue
            free = [c for c in GOOGLE_COLORS if c not in used]
            if not free:
                free = list(GOOGLE_COLORS)
                self.warnings.append(f"More subjects than Google's 11 event colours: "
                                     f"'{code}' ({counts[code]} events) shares a colour.")
            if strategy == "excel" and rgb_of.get(code):
                pick = min(free, key=lambda c: _dist(GOOGLE_COLORS[c], rgb_of[code]))
            else:
                pick = free[0]
            assigned[code] = pick
            used.add(pick)
        assigned[NOTICE_KEY] = None
        for e in events:
            e.color_id = assigned[e.code]
        self.color_map = assigned


def _dist(h1: str, h2: str) -> int:
    a = [int(h1[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(h2[i:i + 2], 16) for i in (1, 3, 5)]
    return sum((x - y) ** 2 for x, y in zip(a, b))
